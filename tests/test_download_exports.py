import csv
import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from resume_shortlisting import config
from resume_shortlisting.excel_writer import to_final_candidate_dataframe, write_final_candidate_sheet
from resume_shortlisting.json_writer import write_json
from resume_shortlisting.models import Candidate, ScoreResult
from resume_shortlisting.profile_exports import profile_path, profiles_zip
from resume_shortlisting.utils import resume_cache_path


class DownloadExportTests(unittest.TestCase):
    def setUp(self):
        self.result = ScoreResult(
            candidate=Candidate(
                name="Download Candidate",
                email="candidate@example.com",
                resume_url="https://example.com/resume.pdf",
            ),
            score=60,
            recommendation="Shortlist",
        )

    def test_xlsx_export_is_non_empty_and_openpyxl_readable(self):
        with tempfile.TemporaryDirectory() as directory:
            output = write_final_candidate_sheet(
                [self.result], Path(directory) / "final_candidate_sheet.xlsx"
            )
            data = output.read_bytes()
            self.assertGreater(len(data), 0)
            workbook = load_workbook(io.BytesIO(data), read_only=True)
            self.assertEqual(workbook.sheetnames, ["Final Candidates"])
            self.assertEqual(workbook.active.max_row, 2)

    def test_json_export_is_non_empty_and_valid(self):
        with tempfile.TemporaryDirectory() as directory:
            output = write_json([self.result], Path(directory) / "report.json")
            data = output.read_bytes()
            self.assertGreater(len(data), 0)
            payload = json.loads(data.decode("utf-8"))
            self.assertEqual(payload["results"][0]["student"], "Download Candidate")

    def test_shortlist_csv_is_valid_utf8_and_contains_candidate(self):
        frame = to_final_candidate_dataframe([self.result])
        data = frame.to_csv(index=False).encode("utf-8-sig")
        self.assertGreater(len(data), 0)
        rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
        self.assertEqual(rows[0]["Student Name"], "Download Candidate")

    def test_selected_profiles_zip_is_non_empty_and_contains_selected_candidate(self):
        original_resumes_dir = config.RESUMES_DIR
        with tempfile.TemporaryDirectory() as directory:
            config.RESUMES_DIR = Path(directory)
            resume_path = resume_cache_path(self.result.candidate.resume_url, self.result.candidate.name, ".pdf")
            resume_path.write_bytes(b"%PDF-1.7 selected resume")
            try:
                data = profiles_zip([self.result])
                self.assertGreater(len(data), 0)
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    names = archive.namelist()
                    self.assertEqual(len(names), 1)
                    self.assertIn("Download_Candidate", names[0])
                    self.assertGreater(len(archive.read(names[0])), 0)
            finally:
                config.RESUMES_DIR = original_resumes_dir


if __name__ == "__main__":
    unittest.main()
