import unittest
import tempfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from resume_shortlisting.models import Candidate
from resume_shortlisting.parser import parse_resume
from resume_shortlisting.skills_kb import load_kb
from resume_shortlisting.talent_pool import GithubEvidence, _deduplicate_students, profile_candidate, score_role, verify_github
from resume_shortlisting.talent_pool import talent_pool_xlsx_bytes, to_dataframe, write_talent_pool_excel
from resume_shortlisting.tech_detector import enrich_projects


class FakeResponse:
    def __init__(self, status, data=None): self.status_code, self._data = status, data or {}
    def json(self): return self._data

class FakeSession:
    def __init__(self, responses): self.responses = responses
    def get(self, url, **kwargs): return self.responses[url]
    def close(self): pass

class TalentPoolTests(unittest.TestCase):
    def setUp(self): self.kb = load_kb()

    def test_normalization_is_boundary_safe(self):
        self.assertEqual(self.kb.canonical_of("React.js"), "React")
        self.assertEqual(self.kb.canonical_of("ReactJS"), "React")
        self.assertEqual(self.kb.canonical_of("NodeJS"), "Node.js")
        self.assertEqual(self.kb.canonical_of("DRF"), "Django REST Framework")
        self.assertFalse(self.kb.skill_in_text("Java", "JavaScript"))

    def test_strong_mern_genai_is_role_specific_and_explainable(self):
        text = """Skills\nReactJS, NodeJS, Express, Mongo DB, JavaScript, RAG, LLM, LangChain, OpenAI\nProjects\nAI Shop\nBuilt a React NodeJS Express Mongo DB e-commerce assistant using RAG, LLM and LangChain.\nGitHub: https://github.com/a/shop\nLive: https://shop.example.com\n"""
        profile = profile_candidate(Candidate(name="A"), text, check_github=False, kb=self.kb)
        self.assertGreater(profile.role_scores["MERN"].score, 50)
        self.assertGreater(profile.role_scores["GenAI"].score, 45)
        self.assertTrue(profile.role_scores["MERN + GenAI"].relevant_projects)
        self.assertLess(profile.role_scores["Java Full Stack"].score, 25)

    def test_projects_are_capped_and_irrelevant_do_not_score(self):
        resume = parse_resume("Skills\nReact, Node.js, Express, MongoDB\nProjects\nOne\nReact Node.js Express MongoDB\nTwo\nReact Node.js Express MongoDB\nThree\nReact Node.js Express MongoDB\nFour\nReact Node.js Express MongoDB")
        enrich_projects(resume.project_list, self.kb)
        score = score_role("MERN", resume, GithubEvidence(), self.kb)
        self.assertLessEqual(score.project_score, 30)
        weak = parse_resume("Projects\nEssay\nA history project")
        enrich_projects(weak.project_list, self.kb)
        self.assertEqual(score_role("MERN", weak, GithubEvidence(), self.kb).project_score, 0)

    def test_github_network_failure_remains_unknown(self):
        session = FakeSession({"https://api.github.com/users/a": FakeResponse(500)})
        evidence = verify_github(["https://github.com/a"], session)
        self.assertEqual(evidence.profile_exists, "UNKNOWN")

    def test_verified_public_repo_readme(self):
        urls = {
            "https://api.github.com/users/a": FakeResponse(200),
            "https://api.github.com/repos/a/repo": FakeResponse(200, {"private": False}),
            "https://api.github.com/repos/a/repo/readme": FakeResponse(200),
            "https://api.github.com/repos/a/repo/commits?per_page=100": FakeResponse(200, []),
        }
        evidence = verify_github(["https://github.com/a/repo"], FakeSession(urls))
        self.assertEqual((evidence.profile_exists, evidence.repository_exists, evidence.repository_public, evidence.readme_exists), ("YES", "YES", "YES", "YES"))

    def test_export_preserves_source_order_then_appends_analysis(self):
        source = {
            "Timestamp": "2026-01-01", "Student UID": "S-1", "Student Name": "A",
            "Native Language": "Telugu", "Which Roles you are Looking for Internship": "MERN",
            "Which technology stack(s) are you proficient in for your internship roles?": "React, Node.js",
            "Share your Updated Resume Drive link (Give Public Access). Ensure your resume includes your latest skills, projects.": "https://example.com/a.pdf",
        }
        profile = profile_candidate(Candidate(name="A", source_data=source), "Skills\nReact, Node.js\nProjects\nApp\nReact Node.js", check_github=False, kb=self.kb)
        frame = to_dataframe([profile])
        self.assertEqual(list(frame.columns[:len(source)]), list(source))
        self.assertEqual(frame.loc[0, "Native Language"], "Telugu")
        self.assertIn("Resume Accessible", frame.columns)
        self.assertIn("MERN Score", frame.columns)

    def test_student_uid_keeps_latest_timestamped_response(self):
        older = Candidate(name="Old", resume_url="https://example.com/old.pdf", source_data={"Student UID": "S-1", "Timestamp": "2026-01-02 09:00:00"})
        latest = Candidate(name="Latest", resume_url="https://example.com/latest.pdf", source_data={"Student UID": "S-1", "Timestamp": "2026-02-02 09:00:00"})
        other = Candidate(name="Other", resume_url="https://example.com/other.pdf", source_data={"Student UID": "S-2", "Timestamp": "2026-01-01"})
        selected = _deduplicate_students([older, latest, other])
        self.assertEqual({candidate.source_data["Student UID"] for candidate in selected}, {"S-1", "S-2"})
        self.assertEqual(next(candidate for candidate in selected if candidate.source_data["Student UID"] == "S-1").name, "Latest")

    def test_multiple_projects_and_complex_source_values_export_to_xlsx(self):
        source = {"student_uid": "S-1", "student_name": "Ana", "email": "ana@example.com", "resume_url": "https://example.com/a.pdf", "additional_notes": [{"source": "form"}], "projects": "AI Career Mentor\x0b; Live Demo; Built career guidance; HealthQueue\x15; Django healthcare platform"}
        text = """Skills\nPython, Django\nProjects\nAI Career Mentor\nTechnologies: Python, Django\nHealthQueue\nTechnologies: Python, Django, SQLite\n"""
        profile = profile_candidate(Candidate(name="Ana", source_data=source), text, check_github=False, kb=self.kb)
        with tempfile.TemporaryDirectory() as directory:
            output = write_talent_pool_excel([profile], Path(directory) / "talent.xlsx")
            workbook = load_workbook(output)
            headers = [cell.value for cell in workbook.active[1]]
            project_value = workbook.active.cell(2, headers.index("projects") + 1).value
        self.assertEqual(workbook.active.max_row, 2)
        self.assertNotIn("\x0b", project_value)
        self.assertNotIn("\x15", project_value)

    def test_download_bytes_are_a_valid_xlsx_workbook(self):
        profile = profile_candidate(Candidate(name="Ana"), "Skills\nPython\nProjects\nPortfolio\nPython", check_github=False, kb=self.kb)
        content = talent_pool_xlsx_bytes([profile])
        self.assertTrue(content.startswith(b"PK"))
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        self.assertIn("Student Talent Pool", workbook.sheetnames)
        self.assertGreaterEqual(workbook["Student Talent Pool"].max_row, 2)


if __name__ == "__main__": unittest.main()
