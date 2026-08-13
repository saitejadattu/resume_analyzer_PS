import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from resume_shortlisting.excel_writer import (
    sanitize_excel_value, to_final_candidate_dataframe, write_excel,
    write_final_candidate_sheet,
)
from resume_shortlisting.models import Candidate, GithubStatus, ScoreResult, SkillMatch


class ExcelSanitizationTests(unittest.TestCase):
    def test_sanitizer_preserves_supported_unicode_and_layout(self):
        value = "✓ Django ₹ café\tline\nnext\x00bad\x15text�"
        self.assertEqual(sanitize_excel_value(value), "✓ Django ₹ café\tline\nnextbadtext�")
        self.assertEqual(sanitize_excel_value(None), "")
        self.assertEqual(sanitize_excel_value(42), 42)

    def test_evidence_with_pdf_control_characters_exports(self):
        evidence = SkillMatch(
            skill="SQL", match_type="project",
            project_name="☑ Python for Everybody \x15 University of Michigan �",
            source="tech_stack", project_github_url="", github_status=GithubStatus.NOT_PROVIDED,
        )
        result = ScoreResult(
            candidate=Candidate(name="Zoë \x00 Candidate", resume_url="https://example.com/a.pdf"),
            score=60, recommendation="Shortlist", keyword_evidence=[evidence],
            matched_skills=["SQL"], matched_in={"SQL": ["Projects"]},
            matched_projects=[evidence.project_name], remarks="Extracted\x1f text",
        )
        with tempfile.TemporaryDirectory() as directory:
            output = write_excel([result], Path(directory) / "safe.xlsx")
            workbook = load_workbook(output)
            values = [cell.value for cell in workbook.active[2]]
        joined = " ".join(str(value) for value in values)
        self.assertIn("☑ Python for Everybody  University of Michigan �", joined)
        self.assertNotIn("\x15", joined)
        self.assertNotIn("\x00", joined)

    def test_final_sheet_preserves_source_columns_and_appends_decision(self):
        result = ScoreResult(
            candidate=Candidate(
                name="Ana", email="ana@example.com", resume_url="https://example.com/a.pdf",
                source_data={"Candidate Name": "Ana", "Phone": "123", "Resume Link": "https://example.com/a.pdf"},
            ),
            score=40, recommendation="Reject", remarks="Rejected because Missing required evidence: Django.",
        )
        frame = to_final_candidate_dataframe([result])
        self.assertEqual(
            list(frame.columns),
            ["Candidate Name", "Phone", "Resume Link", "Score", "Status", "Remarks",
             "Required Keywords Matched", "Required Keywords Missing", "Project Matches",
             "Skills Matches", "Project GitHub Status"],
        )
        self.assertEqual(frame.loc[0, "Status"], "Rejected")
        with tempfile.TemporaryDirectory() as directory:
            output = write_final_candidate_sheet([result], Path(directory) / "final.xlsx")
            self.assertEqual(load_workbook(output).active.max_row, 2)


if __name__ == "__main__":
    unittest.main()
