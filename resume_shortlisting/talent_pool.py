"""Deterministic, independently testable Student Talent Pool pipeline."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse
from zipfile import ZipFile
import requests
import pandas as pd
from openpyxl import load_workbook

from . import config
from .downloader import download_all
from .excel_writer import sanitize_excel_value
from .github_checker import classify_kind
from .models import Candidate, GithubKind, ParsedResume, Project
from .parser import parse_resume
from .role_definitions import COMBINED_TRACKS, ROLE_DEFINITIONS
from .skills_kb import SkillsKB, load_kb
from .tech_detector import enrich_projects
from .text_extractor import extract_text
from .utils import get_logger

logger = get_logger("talent_pool")

@dataclass
class GithubEvidence:
    url: str = ""; profile_exists: str = "UNKNOWN"; repository_exists: str = "UNKNOWN"; repository_public: str = "UNKNOWN"; readme_exists: str = "UNKNOWN"
    commits_last_2_months: int | None = None; commits_current_year: int | None = None; last_commit_date: str = ""

@dataclass
class RoleScore:
    track: str; score: float; tier: str; skill_score: float; project_score: float; github_score: float; evidence_score: float
    matched_required_skills: list[str] = field(default_factory=list); matched_preferred_skills: list[str] = field(default_factory=list); relevant_projects: list[str] = field(default_factory=list)

@dataclass
class TalentProfile:
    candidate: Candidate; status: str = "SUCCESS"; error: str = ""; resume: ParsedResume = field(default_factory=ParsedResume)
    github: GithubEvidence = field(default_factory=GithubEvidence); role_scores: dict[str, RoleScore] = field(default_factory=dict)
    profile_status: str = "Incomplete"; target_ready: str = "NO"; target_priority: str = "D — Low Evidence"; recommended_tracks: list[str] = field(default_factory=list); needs_improvement: list[str] = field(default_factory=list)

def _yes(value: bool) -> str: return "YES" if value else "NO"
def _tier(score: float) -> str: return "A" if score >= 90 else "B" if score >= 75 else "C" if score >= 60 else "D"
def _priority(tier: str) -> str: return {"A":"A — Target First", "B":"B — Strong Candidate", "C":"C — Developing", "D":"D — Low Evidence"}[tier]

def _github_request(url: str, session: requests.Session) -> requests.Response:
    return session.get(url, timeout=config.GITHUB_TIMEOUT, headers={"User-Agent": config.USER_AGENT, **({"Authorization": f"Bearer {config.GITHUB_TOKEN}"} if config.GITHUB_TOKEN else {})})

def _verify_github_url(url: str, session: requests.Session) -> GithubEvidence:
    """Verify one URL; network/rate failures remain UNKNOWN."""
    evidence = GithubEvidence(url=url)
    if not url: return evidence
    try:
        kind, owner, repo = classify_kind(url)
        if kind is GithubKind.UNKNOWN: return evidence
        profile = _github_request(f"https://api.github.com/users/{owner}", session)
        evidence.profile_exists = "YES" if profile.status_code == 200 else "NO" if profile.status_code == 404 else "UNKNOWN"
        if kind is GithubKind.REPOSITORY:
            response = _github_request(f"https://api.github.com/repos/{owner}/{repo}", session)
            if response.status_code == 200:
                data = response.json(); evidence.repository_exists = "YES"; evidence.repository_public = _yes(not data.get("private", True))
                readme = _github_request(f"https://api.github.com/repos/{owner}/{repo}/readme", session)
                evidence.readme_exists = "YES" if readme.status_code == 200 else "NO" if readme.status_code == 404 else "UNKNOWN"
                commits = _github_request(f"https://api.github.com/repos/{owner}/{repo}/commits?per_page=100", session)
                if commits.status_code == 200:
                    dates = [datetime.fromisoformat(x["commit"]["author"]["date"].replace("Z", "+00:00")) for x in commits.json() if x.get("commit", {}).get("author", {}).get("date")]
                    today = datetime.now(timezone.utc); evidence.commits_last_2_months = sum(d >= today - timedelta(days=60) for d in dates); evidence.commits_current_year = sum(d.year == today.year for d in dates); evidence.last_commit_date = max(dates).date().isoformat() if dates else ""
            elif response.status_code == 404: evidence.repository_exists = "NO"; evidence.repository_public = "NO"
    except (requests.RequestException, ValueError, KeyError, TypeError) as exc:
        logger.warning("GitHub verification unavailable for %s: %s", url, exc)
    return evidence

def verify_github(urls: list[str], session: requests.Session | None = None) -> GithubEvidence:
    """Attempt every distinct GitHub URL and retain the strongest evidence.

    A rate limit or connection failure never changes an unknown field to NO.
    """
    unique = list(dict.fromkeys(url for url in urls if url))
    if not unique: return GithubEvidence()
    own = session is None; session = session or requests.Session()
    try:
        checked = [_verify_github_url(url, session) for url in unique]
    finally:
        if own: session.close()
    # Prefer a verified public repository, then any accessible link. This is
    # one row's summary while every discovered URL was still attempted.
    return max(checked, key=lambda item: (item.repository_public == "YES", item.repository_exists == "YES", item.profile_exists == "YES"))

def _detected_skills(resume: ParsedResume, kb: SkillsKB) -> set[str]:
    return set(kb.detect_skills(resume.skills + "\n" + resume.experience + "\n" + resume.projects)) | {t for p in resume.project_list for t in p.technologies}

def _project_match(project: Project, definition: dict[str, list[str]], kb: SkillsKB) -> float:
    techs = set(project.technologies) | set(kb.detect_skills(project.description + "\n" + project.name))
    required = definition["required"]; preferred = definition["preferred"]
    req = sum(s in techs for s in required) / len(required) if required else 0
    pref = sum(s in techs for s in preferred) / len(preferred) if preferred else 0
    return .75 * req + .25 * pref

def score_role(track: str, resume: ParsedResume, github: GithubEvidence, kb: SkillsKB | None = None) -> RoleScore:
    kb = kb or load_kb(); definition = ROLE_DEFINITIONS[track]; skills = _detected_skills(resume, kb)
    required = [s for s in definition["required"] if s in skills]; preferred = [s for s in definition["preferred"] if s in skills]
    skill_score = 25 * len(required) / len(definition["required"]) + (10 * len(preferred) / len(definition["preferred"]) if definition["preferred"] else 0)
    matched = sorted((( _project_match(p, definition, kb), p) for p in resume.project_list), key=lambda item: item[0], reverse=True)
    relevant = [(rel, p) for rel, p in matched if rel >= .25]
    caps = (15, 8, 7); project_score = sum(cap * rel for cap, (rel, _) in zip(caps, relevant[:3])); project_score = min(30, project_score + min(2, sum(bool(p.live_url) for _, p in relevant[:3])))
    github_score = 5 if github.url else 0
    github_score += 5 if "YES" in (github.profile_exists, github.repository_exists) else 0
    github_score += 5 if github.repository_public == "YES" else 0; github_score += 5 if github.readme_exists == "YES" else 0
    evidence_score = (5 if required and resume.skills.strip() else 0) + (5 if any(p.description.strip() for _, p in relevant) else 0) + (5 if any(p.technologies for _, p in relevant) else 0)
    total = round(min(100, skill_score + project_score + github_score + evidence_score), 1)
    return RoleScore(track, total, _tier(total), round(skill_score,1), round(project_score,1), github_score, evidence_score, required, preferred, [p.name for _, p in relevant])

def _combined_score(name: str, left: RoleScore, right: RoleScore, projects: list[Project], kb: SkillsKB) -> RoleScore:
    intersection = any(_project_match(p, ROLE_DEFINITIONS[left.track], kb) >= .25 and _project_match(p, ROLE_DEFINITIONS[right.track], kb) >= .25 for p in projects)
    score = round(min(left.score, right.score) + (min(10, abs(left.score-right.score)/4) if intersection else 0), 1)
    return RoleScore(name, score, _tier(score), left.skill_score, left.project_score, left.github_score, left.evidence_score, relevant_projects=(left.relevant_projects if intersection else []))

def profile_candidate(candidate: Candidate, text: str, *, check_github: bool = True, kb: SkillsKB | None = None) -> TalentProfile:
    kb = kb or load_kb(); resume = parse_resume(text); enrich_projects(resume.project_list, kb)
    github = verify_github(resume.github_urls) if check_github else GithubEvidence(url=next(iter(resume.github_urls), ""))
    scores = {track: score_role(track, resume, github, kb) for track in ROLE_DEFINITIONS}
    for name, left, right in COMBINED_TRACKS: scores[name] = _combined_score(name, scores[left], scores[right], resume.project_list, kb)
    recommended = [n for n, s in scores.items() if s.score >= 60]
    best = max(scores.values(), key=lambda s: s.score); needs = []
    if not github.url: needs.append("Missing GitHub")
    if not resume.project_list: needs.append("No projects")
    if resume.project_list and not any(p.live_url for p in resume.project_list): needs.append("No deployed projects")
    if not any(s.relevant_projects for s in scores.values()): needs.append("No relevant projects")
    status = "Complete" if resume.raw_text and resume.project_list and _detected_skills(resume, kb) else "Incomplete"
    ready = "YES" if best.score >= 75 and status == "Complete" and github.repository_public == "YES" else ("NEEDS REVIEW" if best.score >= 60 else "NO")
    return TalentProfile(candidate, resume=resume, github=github, role_scores=scores, profile_status=status, target_ready=ready, target_priority=_priority(best.tier), recommended_tracks=recommended, needs_improvement=needs)

def _deduplicate_students(candidates: list[Candidate]) -> list[Candidate]:
    """Return one current response per Student UID, keeping the latest one.

    The source sheet is append-only in practice, so repeated UIDs represent a
    student updating their profile/resume. Invalid or absent timestamps retain
    source order as a deterministic fallback.
    """
    selected: dict[str, tuple[pd.Timestamp | None, int, Candidate]] = {}
    no_uid: list[Candidate] = []
    for index, candidate in enumerate(candidates):
        uid = candidate.source_data.get("Student UID", "").strip()
        if not uid:
            no_uid.append(candidate)
            continue
        timestamp = pd.to_datetime(candidate.source_data.get("Timestamp", ""), errors="coerce", utc=True)
        parsed = None if pd.isna(timestamp) else timestamp
        previous = selected.get(uid)
        # Valid timestamps outrank missing/invalid ones; for ties, a later row
        # wins, reflecting the latest submitted response in an exported sheet.
        if previous is None or (parsed is not None and (previous[0] is None or parsed >= previous[0])) or (parsed is None and previous[0] is None and index > previous[1]):
            selected[uid] = (parsed, index, candidate)
    return [item[2] for item in sorted(selected.values(), key=lambda item: item[1])] + no_uid

def run_talent_pool(candidates: list[Candidate], *, check_github: bool = True, settings: config.Settings | None = None, on_progress=None) -> list[TalentProfile]:
    settings = settings or config.DEFAULT_SETTINGS
    unique = _deduplicate_students(candidates)
    downloads = download_all(unique, settings); profiles=[]; kb=load_kb()
    for index, candidate in enumerate(unique, 1):
        item = downloads.get(candidate.resume_url)
        if not item or not item.ok or not item.path: profiles.append(TalentProfile(candidate, status="FAILED", error=item.error if item else "download failed"))
        else:
            try:
                text = extract_text(candidate, item.path)
                if not text:
                    profiles.append(TalentProfile(candidate, status="PARTIAL", error="could not extract resume text"))
                else:
                    profiles.append(profile_candidate(candidate, text, check_github=check_github, kb=kb))
            except Exception as exc: profiles.append(TalentProfile(candidate, status="FAILED", error=f"{type(exc).__name__}: {exc}"))
        logger.info("[%s] Talent pool processing completed: %s", candidate.display_name, profiles[-1].status)
        if on_progress: on_progress(index, len(unique))
    return profiles

def to_dataframe(profiles: list[TalentProfile]) -> pd.DataFrame:
    rows=[]
    for p in profiles:
        # Start with the unmodified source row. Python dict order preserves
        # the spreadsheet's original column order supplied by pandas.
        source=dict(p.candidate.source_data)
        if not source: source={"Student UID":"", "Student Name":p.candidate.display_name, "Email":p.candidate.email}
        row = source | {"Profile Status":p.profile_status, "Target Ready":p.target_ready, "Target Priority":p.target_priority, "Recommended Tracks":", ".join(p.recommended_tracks), "Needs Improvement":", ".join(p.needs_improvement), "Resume Accessible":_yes(bool(p.resume.raw_text)), "Resume Parsed":_yes(bool(p.resume.raw_text)), "Extracted Skills":", ".join(sorted(_detected_skills(p.resume, load_kb()))) if p.resume.raw_text else "", "Extracted Projects":"; ".join(x.name for x in p.resume.project_list), "Processing Status":p.status, "Error":p.error, "Total Projects":len(p.resume.project_list), "GitHub-backed Projects":sum(bool(x.github_url) for x in p.resume.project_list), "Deployed Projects":sum(bool(x.live_url) for x in p.resume.project_list), "GitHub URL":p.github.url, "GitHub Profile Exists":p.github.profile_exists, "GitHub Repository Exists":p.github.repository_exists, "GitHub Repository Public":p.github.repository_public, "GitHub README Exists":p.github.readme_exists, "GitHub Commits Last 2 Months":p.github.commits_last_2_months, "GitHub Commits Current Year":p.github.commits_current_year, "GitHub Last Commit Date":p.github.last_commit_date, "GitHub Verification Status":("Verified" if p.github.repository_public == "YES" else "Unknown" if "UNKNOWN" in (p.github.profile_exists, p.github.repository_exists, p.github.repository_public, p.github.readme_exists) else "Not Verified")}
        row["Relevant Projects"] = max((len(s.relevant_projects) for s in p.role_scores.values()), default=0)
        for name, score in p.role_scores.items(): row[f"{name} Score"] = score.score; row[f"{name} Tier"] = score.tier
        uid = str(source.get("student_uid", source.get("Student UID", "")))
        safe_row: dict[str, object] = {}
        for field, value in row.items():
            try:
                # Handles both complex values and PDF/source control chars.
                safe_row[field] = sanitize_excel_value(value)
            except Exception as exc:  # defensive: one field must not abort a batch
                logger.warning(
                    "[%s] Worksheet serialization fallback for field %s (type=%s, value=%r): %s",
                    uid or p.candidate.display_name, field, type(value).__name__, repr(value)[:500], exc,
                )
                safe_row[field] = sanitize_excel_value(str(value))
        rows.append(safe_row)
    # Final defensive barrier: even a future output column cannot pass a raw
    # container or illegal control character to Streamlit/XLSX consumers.
    frame = pd.DataFrame(rows).map(sanitize_excel_value)
    # Streamlit/Arrow cannot encode an object column containing both ``""``
    # and integers. Use pandas' nullable integer dtype for activity values.
    for column in ("GitHub Commits Last 2 Months", "GitHub Commits Current Year"):
        if column in frame:
            frame[column] = pd.to_numeric(frame[column], errors="coerce").astype("Int64")
    return frame

def talent_pool_xlsx_bytes(profiles: list[TalentProfile]) -> bytes:
    """Build and validate a complete XLSX workbook for Streamlit downloads."""
    frame = to_dataframe(profiles)
    buffer = BytesIO()
    try:
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            frame.to_excel(writer, index=False, sheet_name="Student Talent Pool")
            writer.sheets["Student Talent Pool"].freeze_panes = "A2"
    except Exception:
        logger.exception("Talent Pool XLSX generation failed after worksheet serialization")
        raise
    content = buffer.getvalue()
    try:
        with ZipFile(BytesIO(content)) as archive:
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise ValueError("generated archive is missing required XLSX files")
            if archive.testzip() is not None:
                raise ValueError("generated XLSX archive contains a corrupt entry")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        if "Student Talent Pool" not in workbook.sheetnames:
            raise ValueError("generated workbook is missing the Student Talent Pool sheet")
    except Exception as exc:
        logger.exception("Generated Talent Pool XLSX validation failed")
        raise RuntimeError(f"Generated Talent Pool XLSX is invalid: {exc}") from exc
    logger.info("Generated validated Talent Pool XLSX: %d rows, %d bytes", len(frame), len(content))
    return content


def write_talent_pool_excel(profiles: list[TalentProfile], output: Path) -> Path:
    """Persist the same validated XLSX bytes used by download clients."""
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_bytes(talent_pool_xlsx_bytes(profiles))
    return output
